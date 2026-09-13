"""Belindia, typicality method: the Python side of belindia_typicality.xlsx.

Reads the World Bank 1000-bin zip, the published classes (inversion workbook) and the
survey list (v1 workbook), and writes belindia_typicality_analysis_output.xlsx with
  - the 2025 lines, matrix and per-economy shares (same method as the live sheets)
  - the 1990-2026 history under three yardsticks (year lines, fixed 2025 lines, pooled lines)
  - the robustness table (bandwidth, kernel, giants excluded, survey-only, class-weighted, missing rich)
  - the class curves on the grid for 1990, 2025 and 2026 (for charting)

Method (identical to the workbook): log grid from $0.10 to $2,000 in 0.05 steps; share of an
economy below an edge = count of bins at or below it / 1000; class curve = population-weighted
share per cell / step (each integrates to 1); Gaussian smoothing, bandwidth 0.2; most typical
class = highest smoothed curve where any curve exceeds 0.001; line k = first cell at or beyond
the peak of class k where the most typical class is k+1 or higher, interpolated on the density
difference; shares below a line by MATCH with linear interpolation between bin midpoints.
"""
import subprocess, warnings, numpy as np, pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter as COL
warnings.filterwarnings("ignore")
ZIP = "GlobalDist1000bins_1990_2026_20260324_2021_01_02_PROD_csv.zip"
CSV = "GlobalDist1000bins_1990_2026_20260324_2021_01_02_PROD.csv"
INV = "20260725-inversion-income-classification-2050.xlsx"
BL = "belindia_live.xlsx"
OUT = "belindia_typicality_analysis_output.xlsx"
CL = ["L", "LM", "UM", "H"]
LOG_LO, LOG_HI, STEP, BW, EPS = np.log(0.1), np.log(2000.0), 0.05, 0.2, 1e-3
GRID = np.arange(0, int(round((LOG_HI - LOG_LO) / STEP)) + 1) * STEP + LOG_LO
X = np.exp(GRID); NG = len(GRID); MID = (GRID[:-1] + GRID[1:]) / 2


# ------------------------------------------------------------------ method
def shares_below_grid(W):
    return np.column_stack([(W <= x).sum(1) for x in X]) / 1000.0


def kernel(bw, kind="gauss"):
    if kind == "gauss":
        kk = np.arange(-int(np.ceil(5 * bw / STEP)), int(np.ceil(5 * bw / STEP)) + 1)
        ker = np.exp(-0.5 * (kk * STEP / bw) ** 2)
    else:  # Epanechnikov with the same variance
        r = int(np.ceil(bw / STEP * np.sqrt(5))); kk = np.arange(-r, r + 1)
        ker = np.clip(1 - (kk * STEP / (bw * np.sqrt(5))) ** 2, 0, None)
    return ker / ker.sum()


def class_density(F, cls, pop, bw=BW, kind="gauss", weighted=False):
    ker = kernel(bw, kind); s = {}
    for c in CL:
        m = cls == c
        if m.sum() == 0 or pop[m].sum() == 0:
            s[c] = np.zeros(NG - 1); continue
        h = (np.diff(F[m], axis=1) * pop[m][:, None]).sum(0) / pop[m].sum() / STEP
        s[c] = np.convolve(h, ker, "same") * (pop[m].sum() if weighted else 1)
    return s


def crossings(s):
    D = np.column_stack([s[c] for c in CL]); am = D.argmax(1)
    am = np.where(D.max(1) > EPS, am, -1); peaks = D.argmax(0); out = []
    for k in range(1, 4):
        idx = np.where((am >= k) & (np.arange(len(am)) >= peaks[k - 1]))[0]
        if len(idx) == 0: out.append(np.nan); continue
        j = idx[0]
        if j == 0: out.append(float(np.exp(MID[0]))); continue
        a, b = am[j - 1], am[j]; a = k - 1 if a < 0 else a
        d0 = D[j - 1, a] - D[j - 1, b]; d1 = D[j, a] - D[j, b]
        t = 0.5 if d0 == d1 else min(max(d0 / (d0 - d1), 0.0), 1.0)
        out.append(float(np.exp(MID[j - 1] + t * STEP)))
    return out, am


def share_below(W, z):
    out = np.zeros((len(W), len(z)))
    for i, w in enumerate(W):
        for j, x in enumerate(z):
            m = int((w <= x).sum())
            out[i, j] = (0.0005 * x / w[0] if m == 0 else 1.0 if m == 1000
                         else (m - 0.5 + (x - w[m - 1]) / (w[m] - w[m - 1])) / 1000)
    return out


def matrix(W, cls, pop, z):
    b = share_below(W, z)
    sh = np.column_stack([b[:, 0], b[:, 1] - b[:, 0], b[:, 2] - b[:, 1], 1 - b[:, 2]])
    M = pd.DataFrame(sh * pop[:, None], columns=CL).groupby(cls).sum().reindex(CL).fillna(0)
    return M, sh


def summ(M):
    v = M.values; tot = v.sum(); d = np.trace(v); up = np.triu(v, 1).sum()
    return dict(total=tot, matches=d, poorer=tot - d - up, richer=up)


# ------------------------------------------------------------------ data
def table(ws, hdr):
    rows = list(ws.iter_rows(min_row=hdr, values_only=True))
    return pd.DataFrame([r for r in rows[1:] if r[1]], columns=[str(x) for x in rows[0]])


def load():
    p = subprocess.Popen(["unzip", "-p", ZIP, CSV], stdout=subprocess.PIPE)
    df = pd.concat(pd.read_csv(p.stdout, usecols=["year", "code", "quantile", "welf", "pop"],
                               chunksize=2_000_000), ignore_index=True)
    cl = table(load_workbook(INV, data_only=True)["Classification"], 3).set_index("Code")
    q = load_workbook(BL, data_only=True)["PIP quantiles"]
    survey = {q.cell(r, 2).value: q.cell(r, 4).value for r in range(5, 177)}
    return df, cl, survey


def year_data(df, cl, y, cls_year=None):
    cy = str(cls_year or y); d = df[df.year == y].sort_values(["code", "quantile"])
    W = d.pivot(index="code", columns="quantile", values="welf")
    codes = [c for c in W.index if c in cl.index and cl.loc[c, cy] in CL]
    return codes, W.loc[codes].values.astype(float), cl.loc[codes, cy].values, d.groupby("code")["pop"].sum().loc[codes].values.astype(float)


# ------------------------------------------------------------------ run
def run():
    df, cl, survey = load(); R = {}
    codes, W, cls, p = year_data(df, cl, 2025)
    F = shares_below_grid(W); s = class_density(F, cls, p); z25, _ = crossings(s)
    M, sh = matrix(W, cls, p, z25); R["z25"] = z25; R["M25"] = M
    R["ctry"] = pd.DataFrame(sh, index=codes, columns=CL).assign(cls=cls, pop=p, median=np.median(W, 1), mean=W.mean(1))
    R["dens"] = {2025: s}
    rob = []
    def rec(name, z, W_=W, cls_=cls, p_=p, note=""):
        M_, _ = matrix(W_, cls_, p_, z); S = summ(M_)
        rob.append(dict(variant=name, b1=z[0], b2=z[1], b3=z[2], matches=S["matches"] / S["total"], poorer=S["poorer"] / S["total"], richer=S["richer"] / S["total"], LM_at_L=M_.loc["LM", "L"], note=note))
    rec("Typicality, bandwidth 0.20 (baseline)", z25)
    for bw in [0.05, 0.1, 0.15, 0.3, 0.4]:
        rec(f"Typicality, bandwidth {bw:.2f}", crossings(class_density(F, cls, p, bw))[0])
    rec("Typicality, Epanechnikov kernel (same variance)", crossings(class_density(F, cls, p, kind="epa"))[0])
    for drop, lab in [(["IND"], "India excluded from the LMIC curve"), (["CHN"], "China excluded from the UMIC curve"), (["IND", "CHN"], "India and China excluded from the curves")]:
        m = ~np.isin(codes, drop); rec(f"Typicality, {lab}", crossings(class_density(F[m], cls[m], p[m]))[0], note="everyone still counted")
    m = np.isin(codes, list(survey)); rec("Typicality, curves from survey economies only", crossings(class_density(F[m], cls[m], p[m]))[0], note="everyone still counted")
    rec("Likeliest origin (curves weighted by class population)", crossings(class_density(F, cls, p, weighted=True))[0], note="footnote only: rewards group size")
    for f in [1.5, 2.0]:
        q = np.arange(1, 1001); W2 = W * np.where(q > 900, 1 + (f - 1) * (q - 900) / 100, 1.0)
        rec(f"Missing rich: top decile scaled up to x{f}", crossings(class_density(shares_below_grid(W2), cls, p))[0], W2)
    rec("Threshold translation, mean floors (annex)", [4.74, 11.01, 26.19], note="belindia_live_v2.xlsx V1")
    rec("Threshold translation, median floors (annex)", [3.65, 8.41, 20.81], note="belindia_live_v2.xlsx V2")
    rec("Likeliest-class lines over economy-year means, 2015-2025 (annex)", [4.44, 10.29, 27.87], note="belindia_live_v2.xlsx V7")
    R["rob"] = pd.DataFrame(rob)
    # history and pooled
    Fs, cs, ps, hist = [], [], [], []
    for y in range(1990, 2027):
        c_, W_, cls_, p_ = year_data(df, cl, y); F_ = shares_below_grid(W_); s_ = class_density(F_, cls_, p_); zy, _ = crossings(s_)
        Fs.append(F_); cs.append(cls_); ps.append(p_)
        if y in (1990, 2026): R["dens"][y] = s_
        row = dict(year=y, status="nowcast" if y == 2026 else "actual", pop=p_.sum(), b1=zy[0], b2=zy[1], b3=zy[2])
        for c in CL: row[f"pop_{c}"] = p_[cls_ == c].sum()
        for lab, z in [("year", zy), ("fixed", z25)]:
            S = summ(matrix(W_, cls_, p_, z)[0]); row.update({f"{lab}_{k}": v for k, v in S.items()})
        hist.append(row); R[f"W{y}"] = (c_, W_, cls_, p_)
    zpool, _ = crossings(class_density(np.vstack(Fs), np.concatenate(cs), np.concatenate(ps))); R["zpool"] = zpool
    for i, y in enumerate(range(1990, 2027)):
        c_, W_, cls_, p_ = R[f"W{y}"]; S = summ(matrix(W_, cls_, p_, zpool)[0]); hist[i].update({f"pooled_{k}": v for k, v in S.items()})
    R["hist"] = pd.DataFrame(hist)
    for y in (1990, 2026):
        c_, W_, cls_, p_ = R[f"W{y}"]; R[f"M{y}"] = matrix(W_, cls_, p_, R["hist"].set_index("year").loc[y, ["b1", "b2", "b3"]].values)[0]; R[f"M{y}fixed"] = matrix(W_, cls_, p_, z25)[0]
    c_, W_, cls26, p_ = year_data(df, cl, 2025, 2026); R["Mrelabel"] = matrix(W_, cls26, p_, z25)[0]
    return R


# ------------------------------------------------------------------ Excel
def write(R):
    wb = Workbook(); wb.remove(wb.active); A = lambda **k: Font(name="Arial", size=k.pop("size", 10), **k)
    T, B, N = A(size=13, bold=True), A(bold=True), A()
    def sheet(name, title, note):
        ws = wb.create_sheet(name); ws["A1"] = title; ws["A1"].font = T; ws["A2"] = note; ws["A2"].font = N; return ws
    def frame(ws, df, r0, fm=None, index=False):
        fm = fm or {}; cols = ([df.index.name or ""] if index else []) + list(df.columns)
        for j, h in enumerate(cols, 1): ws.cell(r0, j, h).font = B
        for i, (ix, row) in enumerate(df.iterrows(), r0 + 1):
            vals = ([ix] if index else []) + list(row.values)
            for j, v in enumerate(vals, 1):
                v = None if (isinstance(v, float) and np.isnan(v)) else (v.item() if hasattr(v, "item") else v)
                c = ws.cell(i, j, v); c.font = N
                if fm.get(cols[j - 1]): c.number_format = fm[cols[j - 1]]
        return r0 + len(df)
    rm = sheet("Read me", "Belindia typicality: Python output", __doc__.strip())
    rm.column_dimensions["A"].width = 140
    ws = sheet("Lines and matrix 2025", "2025 typicality lines and matrix (millions)", "Rows: economy class. Columns: standard people live at.")
    for j, v in enumerate(R["z25"], 2): ws.cell(4, j, v).number_format = "0.00"
    ws.cell(4, 1, "Lines b1..b3").font = B
    frame(ws, R["M25"].rename_axis("Class \\ std"), 6, {c: "#,##0" for c in CL}, index=True); S = summ(R["M25"])
    for i, (k, v) in enumerate(S.items()): ws.cell(13 + i, 1, k); ws.cell(13 + i, 2, float(v)).number_format = "#,##0"; ws.cell(13 + i, 3, float(v / S["total"])).number_format = "0.0%"
    for nm, key in [("Matrix 1990 (1990 lines)", "M1990"), ("Matrix 1990 (2025 lines)", "M1990fixed"), ("Matrix 2026 nowcast (2026 lines)", "M2026"), ("Matrix 2026 nowcast (2025 lines)", "M2026fixed"), ("Matrix 2025 with 2026 classes (2025 lines)", "Mrelabel")]:
        w2 = sheet(nm[:31], nm, "Millions of people."); frame(w2, R[key].rename_axis("Class \\ std"), 4, {c: "#,##0" for c in CL}, index=True)
    h = R["hist"].copy()
    for lab in ["year", "fixed", "pooled"]:
        for k in ["matches", "poorer", "richer"]: h[f"{lab}_{k}_share"] = h[f"{lab}_{k}"] / h["pop"]
    hs = sheet("History 1990-2026", "Label fit by year under three yardsticks", "year_* = lines re-estimated each year; fixed_* = 2025 lines; pooled_* = lines from all years pooled. Shares of that year's classified population.")
    frame(hs, h, 4, {**{c: "#,##0" for c in h.columns if c.startswith("pop") or c.endswith(("matches", "poorer", "richer", "total"))}, **{c: "0.0%" for c in h.columns if c.endswith("share")}, "b1": "0.00", "b2": "0.00", "b3": "0.00"})
    rb = sheet("Robustness", "Alternative choices, 2025", "Lines re-derived under each choice; the matrix recomputed."); frame(rb, R["rob"], 4, {"b1": "0.00", "b2": "0.00", "b3": "0.00", "matches": "0.0%", "poorer": "0.0%", "richer": "0.0%", "LM_at_L": "#,##0"}); rb.column_dimensions["A"].width = 58
    ct = sheet("Economies 2025", "Share of each economy's people at each standard, 2025", "L, LM, UM, H = share at each standard; cls = own class."); frame(ct, R["ctry"].rename_axis("code"), 4, {c: "0.0%" for c in CL} | {"pop": "#,##0.0", "median": "0.00", "mean": "0.00"}, index=True)
    for y, s in R["dens"].items():
        d = pd.DataFrame({"log_mid": MID, "welfare": np.exp(MID)} | {c: s[c] for c in CL})
        de = sheet(f"Curves {y}", f"Class curves on the grid, {y}", "Smoothed density per unit of log welfare; plot against welfare on a log axis."); frame(de, d, 4, {"welfare": "0.00", "log_mid": "0.000"} | {c: "0.0000" for c in CL})
    wb.save(OUT); print("wrote", OUT)


if __name__ == "__main__":
    R = run(); write(R)
    import pickle; pickle.dump(R, open("typ_script_results.pkl", "wb"))
